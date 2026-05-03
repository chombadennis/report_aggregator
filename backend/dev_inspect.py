from docx import Document
import sys
import os
import re

def normalize(s):
    return re.sub(r'[^A-Z0-9]', '', s.upper())

def inspect_docx(path):
    if not os.path.exists(path):
        print(f"Error: File not found at {path}")
        return

    doc = Document(path)
    print(f"\n{'='*80}")
    print(f"🔍 INSPECTING: {os.path.basename(path)}")
    print(f"Location: {path}")
    print(f"Structure: {len(doc.paragraphs)} Paragraphs, {len(doc.tables)} Tables")
    print(f"{'='*80}\n")

    # 1. Key Section Search
    keywords = ["LABOUR", "MATERIAL", "WEATHER", "PROGRESS", "INSTRUCTION", "CUBE"]
    print("📍 SECTION HEADERS FOUND:")
    for i, para in enumerate(doc.paragraphs):
        text = para.text.strip()
        if text and any(k in text.upper() for k in keywords):
            print(f"  [Para {i:3}] {text[:100]}")

    # 2. Table Deep Dive
    print("\n📊 TABLE STRUCTURES:")
    for i, table in enumerate(doc.tables):
        print(f"\n--- TABLE {i} ({len(table.rows)} rows x {len(table.columns)} columns) ---")
        # Try to identify the table by its first row content
        try:
            header_cells = [c.text.strip() for c in table.rows[0].cells]
            print(f"  Header: {header_cells}")
            
            # Show a few sample rows
            sample_limit = 3
            for r_idx, row in enumerate(table.rows[1:sample_limit+1], 1):
                cells = [c.text.strip()[:30] for c in row.cells]
                print(f"  Row {r_idx}: {cells}")
        except Exception as e:
            print(f"  [Error reading table: {e}]")

def main():
    if len(sys.argv) < 2:
        print("Usage: python dev_inspect.py <path_to_docx>")
        print("Example: python dev_inspect.py weekly_template.docx")
        return

    target = sys.argv[1]
    inspect_docx(target)

if __name__ == "__main__":
    main()
