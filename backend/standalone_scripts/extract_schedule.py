import os
import json
import openpyxl
from datetime import datetime

def parse_excel_schedule(filepath):
    """
    Parses a schedule Excel file and builds a hierarchical JSON structure 
    based on the cell indentation level. Handles dynamic depth.
    """
    wb = openpyxl.load_workbook(filepath, data_only=True)
    sheet = wb.active
    
    # Find columns dynamically based on header
    header_row = list(sheet.iter_rows(min_row=1, max_row=1))[0]
    headers = [str(cell.value).strip().lower() if cell.value else "" for cell in header_row]
    
    # Defaults for older excel files without proper headers or different columns
    task_col_idx = 0
    start_col_idx = 1
    finish_col_idx = 2
    
    if 'task name' in headers:
        task_col_idx = headers.index('task name')
    if 'start date' in headers:
        start_col_idx = headers.index('start date')
    if 'finish date' in headers:
        finish_col_idx = headers.index('finish date')

    rows = list(sheet.iter_rows(min_row=2))
    
    root_tasks = []
    stack = []  # will hold tuples of (node_dict, indent_level)
    
    for row in rows:
        if len(row) <= max(task_col_idx, start_col_idx, finish_col_idx):
            continue

        task_cell = row[task_col_idx]
        start_cell = row[start_col_idx]
        finish_cell = row[finish_col_idx]
        
        if not task_cell.value:
            continue
            
        task_name = str(task_cell.value).strip()
        indent = task_cell.alignment.indent if task_cell.alignment.indent else 0
        
        start_date = None
        if isinstance(start_cell.value, datetime):
            start_date = start_cell.value.strftime('%Y-%m-%d')
        elif start_cell.value:
            start_date = str(start_cell.value).strip()
            
        finish_date = None
        if isinstance(finish_cell.value, datetime):
            finish_date = finish_cell.value.strftime('%Y-%m-%d')
        elif finish_cell.value:
            finish_date = str(finish_cell.value).strip()
        
        node = {
            "name": task_name,
            "start_date": start_date,
            "finish_date": finish_date,
            "children": []
        }
        
        # Pop from stack until we find the proper parent
        # A proper parent must have an indent strictly less than the current indent
        while stack and stack[-1][1] >= indent:
            stack.pop()
            
        if not stack:
            # If stack is empty, this is a top-level node
            root_tasks.append(node)
        else:
            # Otherwise, add to the children of the last node in the stack
            stack[-1][0]["children"].append(node)
            
        stack.append((node, indent))
        
    return root_tasks

def process_all_excel_files(directory, output_dir):
    if not os.path.exists(output_dir):
        os.makedirs(output_dir)
        
    for filename in os.listdir(directory):
        if filename.startswith('Makindu_AHP_') and filename.endswith('.xlsx'):
            filepath = os.path.join(directory, filename)
            print(f"Extracting: {filename}")
            
            try:
                hierarchy = parse_excel_schedule(filepath)
                
                # Create output filename
                base_name = filename.replace('.xlsx', '')
                out_path = os.path.join(output_dir, f"{base_name}.json")
                
                with open(out_path, 'w') as f:
                    json.dump(hierarchy, f, indent=2)
                    
                print(f"Saved to: {out_path}")
            except Exception as e:
                print(f"Failed to process {filename}: {e}")

if __name__ == "__main__":
    # Point this to your maks_ahp root folder where the excel files are
    script_dir = os.path.dirname(__file__)
    root_dir = os.path.abspath(os.path.join(script_dir, '..', '..', '..'))
    output_dir = os.path.abspath(os.path.join(script_dir, '..', 'extracted_schedules'))
    process_all_excel_files(root_dir, output_dir)
