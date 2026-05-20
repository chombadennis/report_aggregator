from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import os
from datetime import datetime
import json

class AuditReportGenerator:
    def __init__(self):
        pass

    def generate_report(self, output_path, data, insights, financials=None):
        doc = Document()
        
        # --- COVER PAGE ---
        self._add_cover_page(doc, data)
        
        # --- EXECUTIVE SUMMARY ---
        doc.add_page_break()
        self._add_heading(doc, "1.0 EXECUTIVE SUMMARY", level=1)
        doc.add_paragraph(insights.get("executive_summary", "No summary available."))
        
        # --- RECALIBRATION EXECUTIVE AUDIT (VERBATIM REQUIREMENT) ---
        self._add_heading(doc, "1.1 RECALIBRATION EXECUTIVE AUDIT", level=2)
        
        # Extract months with fallback to last two records
        m_list = financials.get("monthly_financials", []) if financials else []
        completed_month = None
        ongoing_month = None
        
        if len(m_list) >= 2:
            # Usually the last one is ongoing, the one before is completed
            ongoing_month = m_list[-1]
            completed_month = m_list[-2]
        elif len(m_list) == 1:
            ongoing_month = m_list[0]
            
        latest_week = financials.get("weekly_financials", [{}])[-1] if financials else {}
        variance = abs(latest_week.get("variance", 0))
        
        # Dynamically compute the last day of the ongoing month (no hardcoded "31st")
        def _get_month_last_day(month_key: str) -> str:
            """Returns the last day of a month as an ordinal string e.g. '31st', '30th', '28th'."""
            import calendar
            try:
                dt = datetime.strptime(month_key, "%B %Y")
                last_day = calendar.monthrange(dt.year, dt.month)[1]
                if last_day in (11, 12, 13):
                    suffix = "th"
                elif last_day % 10 == 1:
                    suffix = "st"
                elif last_day % 10 == 2:
                    suffix = "nd"
                elif last_day % 10 == 3:
                    suffix = "rd"
                else:
                    suffix = "th"
                return f"{last_day}{suffix}"
            except Exception:
                return "31st"  # safe fallback
        
        ongoing_month_last_day = _get_month_last_day(ongoing_month.get("month", "Current Month")) if ongoing_month else "31st"
        
        p = doc.add_paragraph()
        if completed_month and ongoing_month:
            p.add_run(f"Analysis of the last completed month ({completed_month.get('month', 'Previous Month')}) shows the project achieved {completed_month.get('actual_production', 0):.2f}% production against an envisaged S-curve target of {completed_month.get('envisaged_production', 0):.2f}%. ").bold = False
            
            p.add_run(f"As of the latest live report in mid-{ongoing_month.get('month', 'Current Month').split(' ')[0]}, the cumulative variance has widened to {variance:.2f}% behind the project baseline S-curve. ").bold = True
            
            p.add_run(f"To hit the newly recalibrated milestone of {ongoing_month.get('target_rolling_month_end', 0):.2f}% by {ongoing_month.get('month', 'Current Month')} {ongoing_month_last_day}, the contractor must maintain a strict velocity of {ongoing_month.get('required_weekly', 0):.2f}% per week for the remainder of {ongoing_month.get('month', 'Current Month')}.")
        else:
            p.add_run("Recalibration data is still initializing for the current period. Baseline linear tracking remains at 0.96% per week.")

        # P.S. Section
        doc.add_paragraph("")
        ps_p = doc.add_paragraph()
        run = ps_p.add_run("P.S. Mathematical Logic & S-Curve Forgiveness:")
        run.bold = True
        run.underline = True
        
        pct_time = latest_week.get("pct_time", 0)
        pct_work = latest_week.get("pct_work", 0)
        envisaged_pct_work = latest_week.get("envisaged_pct_work", 0)
        slippage_gap = latest_week.get("slippage_gap", 0)
        req_weekly = ongoing_month.get('required_weekly', 0) if ongoing_month else 0.96
        m_name = ongoing_month.get('month', 'Current Month') if ongoing_month else 'Current Month'
        
        doc.add_paragraph(
            f"The {variance:.2f}% variance represents the true cumulative S-curve progress deficit. The S-curve expected progress to be at {envisaged_pct_work:.2f}%, but actual progress is {pct_work:.2f}%. "
            f"This is fundamentally distinct from the pure calendar Slippage Gap of {slippage_gap:.2f}% (which is the elapsed project time of {pct_time:.2f}% minus work completed). "
            f"If the system used a straight linear mathematical baseline, the contractor would be heavily penalized for the naturally slow site mobilization phase, and the deficit would incorrectly match the massive {slippage_gap:.2f}% slippage gap. "
            f"Instead, the mathematical S-Curve mathematically forgives the slow start. It calculates that the project was only ever expected to be at {envisaged_pct_work:.2f}% by this date. "
            f"This true, realistic {variance:.2f}% backlog is the exact mathematical deficit that forced the required weekly velocity to shift from the original baseline up to the current {req_weekly:.2f}% in order to recover the timeline."
        )
        
        # --- MONTHLY PRODUCTION CALIBRATION ---
        if financials and financials.get("monthly_financials"):
            doc.add_page_break()
            self._add_heading(doc, "2.0 MONTHLY PRODUCTION CALIBRATION", level=1)
            
            # Monthly Table
            m_table = doc.add_table(rows=1, cols=6)
            m_table.style = 'Table Grid'
            hdr = m_table.rows[0].cells
            for idx, text in enumerate(["Month", "Start %", "End %", "Actual (x)", "Envisaged (y)", "Variance (k)"]):
                hdr[idx].text = text
                hdr[idx].paragraphs[0].runs[0].bold = True
            
            for m in financials["monthly_financials"]:
                # ONLY show completed months in the historical comparison table
                if m.get("is_ongoing"):
                    continue
                    
                row = m_table.add_row().cells
                row[0].text = m["month"]
                row[1].text = f"{m['start_pct']:.2f}%"
                row[2].text = f"{m['end_pct']:.2f}%"
                row[3].text = f"{m['actual_production']:.2f}%"
                row[4].text = f"{m['envisaged_production']:.2f}%"
                row[5].text = f"{m['variance']:+.2f}%"

            # Add ongoing target note
            ongoing = financials["monthly_financials"][-1]
            if ongoing.get("is_ongoing"):
                p = doc.add_paragraph()
                run = p.add_run(f"\nONGOING CALIBRATION ({ongoing['month']}):")
                run.bold = True
                doc.add_paragraph(f"The project is currently tracking at {ongoing['end_pct']:.2f}% progress.")
                doc.add_paragraph(f"- Baseline Month Target: {ongoing.get('target_fixed_month_end', 0):.2f}% (Production: {ongoing.get('production_planned_fixed', 0):.2f}%)")
                doc.add_paragraph(f"- Recalibrated Rolling Target: {ongoing.get('target_rolling_month_end', 0):.2f}% (Production: {ongoing.get('production_required_rolling', 0):.2f}%)")
                doc.add_paragraph(f"Required Velocity: {ongoing['required_weekly']:.2f}% per week.")

        # --- RECALIBRATION CHAIN ---
        if financials and financials.get("weekly_financials"):
            self._add_heading(doc, "3.0 PRODUCTION RECALIBRATION CHAIN (WEEKLY AUDIT)", level=1)
            doc.add_paragraph("Tactical weekly performance variance (k) based on dynamic recalibration.")
            
            w_table = doc.add_table(rows=1, cols=5)
            w_table.style = 'Table Grid'
            hdr = w_table.rows[0].cells
            for idx, text in enumerate(["Reporting Week", "Actual (x)", "Envisaged (y)", "Variance (k)", "Next Target"]):
                hdr[idx].text = text
                hdr[idx].paragraphs[0].runs[0].bold = True
                
            for w in financials["weekly_financials"]:
                row = w_table.add_row().cells
                row[0].text = w["label"]
                row[1].text = f"{w['weekly_actual']:.2f}%"
                row[2].text = f"{w['weekly_envisaged']:.2f}%"
                row[3].text = f"{w['weekly_variance']:+.2f}%"
                row[4].text = f"{w['required_future_rate']:.2f}%"

        # --- SWOT ANALYSIS ---
        doc.add_page_break()
        self._add_heading(doc, "4.0 SWOT & PERFORMANCE INTELLIGENCE", level=1)
        
        swot = insights.get("swot", {})
        sections = [
            ("STRENGTHS", swot.get("strengths", []), "008000"), # Green
            ("WEAKNESSES", swot.get("weaknesses", []), "FF0000"), # Red
            ("OPPORTUNITIES", swot.get("opportunities", []), "FFA500"), # Orange
            ("THREATS", swot.get("threats", []), "0000FF") # Blue
        ]
        
        for title, items, color in sections:
            p = doc.add_paragraph()
            run = p.add_run(title)
            run.bold = True
            run.font.color.rgb = RGBColor.from_string(color)
            
            if not items:
                doc.add_paragraph("No items identified.")
            else:
                for item in items:
                    doc.add_paragraph(item, style='List Bullet')
            doc.add_paragraph("")

        # --- STRATEGIC RECOMMENDATIONS ---
        doc.add_page_break()
        self._add_heading(doc, "5.0 STRATEGIC RECOMMENDATIONS", level=1)
        
        recs = insights.get("recommendations", {})
        
        self._add_heading(doc, "5.1 TO THE CLIENT (PM)", level=2)
        for item in recs.get("to_client", ["No recommendations available."]):
            doc.add_paragraph(item, style='List Bullet')
            
        self._add_heading(doc, "5.2 TO THE CONTRACTOR", level=2)
        for item in recs.get("to_contractor", ["No recommendations available."]):
            doc.add_paragraph(item, style='List Bullet')

        # --- PROJECT CORRESPONDENCE & KEY COMMUNICATIONS ---
        doc.add_page_break()
        self._add_heading(doc, "6.0 PROJECT CORRESPONDENCE & COMMUNICATIONS", level=1)
        doc.add_paragraph(
            "The following section lists the formal project correspondence, contractor claims, "
            "and client/project manager instructions uploaded and analyzed by the Field Intelligence system."
        )

        docs_dir = os.path.join("cache", "project_documents")
        docs_list = []
        if os.path.exists(docs_dir):
            for f_name in os.listdir(docs_dir):
                if f_name.endswith(".json"):
                    try:
                        with open(os.path.join(docs_dir, f_name), "r") as f:
                            docs_list.append(json.load(f))
                    except:
                        continue

        if not docs_list:
            doc.add_paragraph("No correspondence or communications uploaded for this audit period.")
        else:
            # Sort by date sent descending
            docs_list.sort(key=lambda x: x.get("date_sent", ""), reverse=True)
            
            # Create a 2-column table with 4 rows for each correspondence item
            for item in docs_list:
                table = doc.add_table(rows=4, cols=2)
                table.style = 'Table Grid'
                
                # Extract and format date/category
                cat_label = str(item.get("category", "")).upper()
                date_sent = item.get("date_sent", "N/A")
                date_val = f"{date_sent}\n[{cat_label}]"
                
                # Title
                title_val = item.get("title", "Untitled")
                
                # Sender & Recipient
                sender = item.get("sender", "N/A")
                recipient = item.get("recipient", "N/A")
                sender_val = f"From: {sender}\nTo: {recipient}"
                
                # AI Summary & Implications
                ai_a = item.get("ai_analysis", {})
                implications = "None"
                summary = item.get("summary", "")
                if isinstance(ai_a, dict):
                    implications = ai_a.get("contractual_implications", "None")
                summary_val = f"Summary:\n{summary}\n\nContractual Implications:\n{implications}"
                
                # Row data mapping
                row_data = [
                    ("Date / Type", date_val),
                    ("Document Title", title_val),
                    ("Sender → Recipient", sender_val),
                    ("AI Summary & Contractual Implications", summary_val)
                ]
                
                # Populate cells
                for r_idx, (hdr_text, val_text) in enumerate(row_data):
                    cells = table.rows[r_idx].cells
                    
                    # Header column (Left side)
                    cells[0].text = hdr_text
                    cells[0].paragraphs[0].runs[0].bold = True
                    
                    # Value column (Right side)
                    cells[1].text = val_text
                    
                    # Style title value bold
                    if r_idx == 1:
                        cells[1].paragraphs[0].runs[0].bold = True
                
                # Set cell widths explicitly
                for row in table.rows:
                    row.cells[0].width = Inches(2.0)
                    row.cells[1].width = Inches(4.5)
                
                # Apply custom styling (margins, shading, custom borders)
                self._apply_custom_table_styles(table)
                
                # Space between tables
                doc.add_paragraph("")
        # --- CLOSING ---
        doc.add_paragraph("\n\nReport generated by Construction Aggregator.")
        doc.add_paragraph(f"Date: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")

        doc.save(output_path)
        return output_path

    def _add_cover_page(self, doc, data):
        # Center alignment for cover
        for i in range(10): doc.add_paragraph() # Spacer
        
        title = doc.add_paragraph("PROJECT AUDIT & RISK INTELLIGENCE REPORT")
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = title.runs[0]
        run.bold = True
        run.font.size = Pt(24)
        run.font.color.rgb = RGBColor(47, 84, 150) # Dark Blue
        
        subtitle = doc.add_paragraph("Tactical Site Performance Analysis")
        subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
        subtitle.runs[0].font.size = Pt(14)
        
        for i in range(5): doc.add_paragraph() # Spacer
        
        info = doc.add_paragraph(f"Generated for: Makindu Affordable Housing Project")
        info.alignment = WD_ALIGN_PARAGRAPH.CENTER
        
        date_para = doc.add_paragraph(f"Report Date: {datetime.now().strftime('%B %d, %Y')}")
        date_para.alignment = WD_ALIGN_PARAGRAPH.CENTER

    def _add_heading(self, doc, text, level=1):
        h = doc.add_heading(text, level=level)
        # Custom styling could be added here

    def _calculate_slippage(self, trend):
        try:
            time = float(trend.get('time_progress', 0))
            work = float(trend.get('financial_progress', 0))
            return round(time - work, 2)
        except:
            return 0

    def _apply_custom_table_styles(self, table):
        """
        Applies professional styling to the 2-column table:
        - Sets custom light grey thin borders.
        - Sets standard cell margins (padding) for professional spacing.
        - Adds beautiful light blue-gray shading to the first column (header column).
        - Applies a nice thick double-line border on the bottom of the last row.
        """
        for r_idx, row in enumerate(table.rows):
            for c_idx, cell in enumerate(row.cells):
                tcPr = cell._tc.get_or_add_tcPr()
                
                # Set cell padding (margins) for a clean professional look (6pt top/bottom, 7.5pt left/right)
                tcMar = tcPr.find(qn('w:tcMar'))
                if tcMar is None:
                    tcMar = OxmlElement('w:tcMar')
                    tcPr.append(tcMar)
                for edge, size in [('top', '120'), ('bottom', '120'), ('left', '150'), ('right', '150')]:
                    elem = tcMar.find(qn(f'w:{edge}'))
                    if elem is None:
                        elem = OxmlElement(f'w:{edge}')
                        tcMar.append(elem)
                    elem.set(qn('w:w'), size)
                    elem.set(qn('w:type'), 'dxa')
                
                # Shade the left header column
                if c_idx == 0:
                    shd = OxmlElement('w:shd')
                    shd.set(qn('w:val'), 'clear')
                    shd.set(qn('w:color'), 'auto')
                    shd.set(qn('w:fill'), 'F2F4F8')  # Very light blue-gray shading
                    tcPr.append(shd)
                
                # Customize borders
                tcBorders = tcPr.first_child_found_in("w:tcBorders")
                if tcBorders is None:
                    tcBorders = OxmlElement('w:tcBorders')
                    tcPr.append(tcBorders)
                
                # Define border configurations
                borders = {
                    'top': {'val': 'single', 'sz': '4', 'space': '0', 'color': 'D3D3D3'},
                    'left': {'val': 'single', 'sz': '4', 'space': '0', 'color': 'D3D3D3'},
                    'right': {'val': 'single', 'sz': '4', 'space': '0', 'color': 'D3D3D3'},
                    'bottom': {'val': 'single', 'sz': '4', 'space': '0', 'color': 'D3D3D3'}
                }
                
                # Use a double border on the bottom of the last row to separate items
                if r_idx == len(table.rows) - 1:
                    borders['bottom'] = {'val': 'double', 'sz': '12', 'space': '0', 'color': '2F5496'}
                
                for edge, attr in borders.items():
                    tag = f'w:{edge}'
                    elem = tcBorders.find(qn(tag))
                    if elem is None:
                        elem = OxmlElement(tag)
                        tcBorders.append(elem)
                    for k, v in attr.items():
                        elem.set(qn(f'w:{k}'), str(v))
