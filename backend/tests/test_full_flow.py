import asyncio
import os
import sys

# Add the backend directory to sys.path so we can import our modules
sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), '..')))

from parser import ReportParser
from aggregator import Aggregator
from generator import ReportGenerator

def normalize_formatting(path):
    import docx
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    
    print(f"Normalizing layout formatting for {os.path.basename(path)}...")
    doc = Document(path)
    
    # 1. Normalize tables 3 to 10 only (these are the ones from Section E to Weather)
    for idx, table in enumerate(doc.tables):
        if 3 <= idx <= 10:
            table.alignment = WD_TABLE_ALIGNMENT.LEFT
            
            tblPr = table._tbl.tblPr
            tblInds = tblPr.xpath('w:tblInd')
            if tblInds:
                for tblInd in tblInds:
                    tblPr.remove(tblInd)
            
            new_tblInd = OxmlElement('w:tblInd')
            new_tblInd.set(qn('w:w'), '0')
            new_tblInd.set(qn('w:type'), 'dxa')
            tblPr.append(new_tblInd)
            
    # 2. Normalize paragraphs starting from Section E onwards
    # Find the start paragraph of Section E dynamically
    start_normalize_idx = 33 # default fallback
    for idx, para in enumerate(doc.paragraphs):
        text = para.text.upper().strip()
        if "E." in text and "PROGRESS" in text:
            start_normalize_idx = idx
            break
            
    print(f"Normalizing paragraphs starting from index {start_normalize_idx}...")
    for idx, para in enumerate(doc.paragraphs):
        if idx >= start_normalize_idx:
            pPr = para._p.get_or_add_pPr()
            ind_elems = pPr.xpath('w:ind')
            if ind_elems:
                for ind in ind_elems:
                    pPr.remove(ind)
                    
            para.paragraph_format.left_indent = None
            para.paragraph_format.first_line_indent = None
            para.paragraph_format.right_indent = None
                
    doc.save(path)
    print(f"[OK] Document formatting normalized successfully.")

async def main():
    # ── Configuration ──
    REPORTS_DIR = r"../../dailies"
    DAILY_PDFS = [
        "MAKINDU AHP DAILY PROGRESS REPORT Monday  13th April 2026-1.pdf",
        "MAKINDU AHP DAILY PROGRESS REPORT Tuesday  14th April 2026-1.pdf",
        "MAKINDU AHP DAILY PROGRESS REPORT Wednesday 15th April 2026.pdf",
        "MAKINDU AHP DAILY PROGRESS REPORT Thursday 16th April 2026.pdf",
        "MAKINDU AHP DAILY PROGRESS REPORT Friday 17th April 2026.pdf",
        "MAKINDU AHP DAILY PROGRESS REPORT Saturday 18th April 2026.pdf",
        "MAKINDU AHP DAILY PROGRESS REPORT Sunday 19th April 2026.pdf",
    ]
    template_path = "weekly_template.docx"
    output_path = os.path.join("tests", "tests_output", "FULL_WEEK_TEST_REPORT.docx")
    session_dir = "full_flow_session"
    os.makedirs(os.path.dirname(output_path), exist_ok=True)

    # 1. Initialize
    parser = ReportParser(cache_dir="cache")
    aggregator = Aggregator()
    generator = ReportGenerator(template_path)

    # 2. Parse all 7 days in parallel
    pdfs = [os.path.join(REPORTS_DIR, f) for f in DAILY_PDFS]
    print(f"--- STEP 1: PARSING {len(pdfs)} DAILY REPORTS ---")
    tasks = [parser.parse_report(pdf, session_dir, "DAILY") for pdf in pdfs]
    daily_results = await asyncio.gather(*tasks)

    # 3. Aggregate into Weekly Summary
    print("\n--- STEP 2: AGGREGATING INTO WEEKLY SUMMARY ---")
    weekly_data = await aggregator.compile_weekly_data(daily_results)
    
    # Add project metadata for the cover page
    weekly_data.update({
        "title": "MAKINDU AHP WEEKLY PROGRESS REPORT",
        "time_elapsed": "20 Weeks",
        "pct_period": "20%",
        "pct_work": "15%",
        "report_date": "13th - 19th April 2026"
    })

    # 4. Generate the Word Document
    print(f"\n--- STEP 3: GENERATING FINAL DOCX ---")
    print(f"Injecting data into {template_path}...")
    generator.generate_report(output_path, weekly_data, "WEEKLY")
    
    # 5. Normalize layout formatting (indentation and alignment)
    print(f"\n--- STEP 4: NORMALIZING FORMATTING (TEST ONLY) ---")
    normalize_formatting(output_path)

    print(f"\n[OK] SUCCESS! Open your report: {os.path.abspath(output_path)}")

if __name__ == "__main__":
    asyncio.run(main())
