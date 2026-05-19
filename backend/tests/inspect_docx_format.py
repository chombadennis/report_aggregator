import docx
from docx import Document
from docx.shared import Inches
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
import os

def inspect_format(path):
    print(f"\n==========================================")
    print(f"File: {os.path.basename(path)}")
    print(f"==========================================")
    
    doc = Document(path)
    
    # Inspect tables
    print("\n--- TABLES ---")
    for idx, table in enumerate(doc.tables):
        alignment = table.alignment
        # Check tblPr properties using XML
        tblPr = table._tbl.tblPr
        tblInd = tblPr.xpath('w:tblInd')
        tblInd_val = tblInd[0].get(docx.oxml.ns.qn('w:w')) if tblInd else "None"
        
        # Read first cell text to identify table
        first_cell = table.rows[0].cells[0].text.strip().replace("\n", " ") if table.rows else ""
        print(f"Table {idx:2d}: {first_cell[:40]:40s} | Alignment: {alignment} | tblInd: {tblInd_val}")
        
    # Inspect some paragraphs
    print("\n--- SAMPLE PARAGRAPHS & SECTIONS ---")
    headers = ["HEALTH AND SAFETY", "SECURITY", "VISITORS", "CHALLENGES", "WEEK", "PROGRESS REPORT"]
    for idx, para in enumerate(doc.paragraphs):
        text = para.text.strip()
        if any(h in text.upper() for h in headers) or (len(text) > 0 and idx < 30):
            pPr = para._p.get_or_add_pPr()
            ind = pPr.xpath('w:ind')
            ind_left = ind[0].get(docx.oxml.ns.qn('w:left')) if ind else "None"
            ind_first_line = ind[0].get(docx.oxml.ns.qn('w:firstLine')) if ind else "None"
            
            jc = pPr.xpath('w:jc')
            jc_val = jc[0].get(docx.oxml.ns.qn('w:val')) if jc else "None"
            
            style_name = para.style.name if para.style else "No Style"
            ind_left = str(ind_left)
            ind_first_line = str(ind_first_line)
            jc_val = str(jc_val)
            style_name = str(style_name)
            
            print(f"Para {idx:3d}: {text[:30]:30s} | Style: {style_name:15s} | LeftInd: {ind_left:6s} | FirstLine: {ind_first_line:6s} | Jc (Align): {jc_val}")

if __name__ == "__main__":
    template_path = "weekly_template.docx"
    output_path = "tests/tests_output/FULL_WEEK_TEST_REPORT.docx"
    
    if os.path.exists(template_path):
        inspect_format(template_path)
    else:
        print(f"Template path not found: {template_path}")
    if os.path.exists(output_path):
        inspect_format(output_path)
    else:
        print(f"Output path not found: {output_path}")
