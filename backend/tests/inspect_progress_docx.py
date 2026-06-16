import os
import sys
from docx import Document

def main():
    path = "tests/tests_output/PROGRESS_TEST_REPORT.docx"
    if not os.path.exists(path):
        print("Report not found!")
        return

    doc = Document(path)
    print("Document loaded successfully.")
    
    # Check sections
    print(f"Number of sections: {len(doc.sections)}")
    for idx, sec in enumerate(doc.sections):
        print(f"Section {idx}: different_first_page_header_footer = {sec.different_first_page_header_footer}")
        
        # Check header
        header = sec.header
        print(f"  Header paragraphs count: {len(header.paragraphs)}")
        for i, p in enumerate(header.paragraphs):
            print(f"    Paragraph {i}: text = '{p.text}', alignment = {p.alignment}")
            runs_with_pics = [r for r in p.runs if 'pic' in r._element.xml]
            print(f"      Contains pictures: {len(runs_with_pics) > 0}")

    # Check cover page paragraphs
    print("\nCover Page Paragraphs:")
    for idx, para in enumerate(doc.paragraphs[:10]):
        text = para.text.strip()
        has_pic = any('pic' in r._element.xml for r in para.runs)
        print(f"  Para {idx:2d}: '{text}' | Has pic: {has_pic}")

    # Check last few paragraphs for signature
    print("\nLast Paragraphs (Signature):")
    for idx, para in enumerate(doc.paragraphs[-10:]):
        text = para.text.strip()
        print(f"  Para {idx:2d}: '{text}'")

if __name__ == "__main__":
    main()
