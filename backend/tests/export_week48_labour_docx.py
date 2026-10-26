import sys, os, json, glob, asyncio, re
from datetime import datetime, timedelta
sys.path.insert(0, os.path.join(os.path.dirname(__file__), '..'))
os.chdir(os.path.join(os.path.dirname(__file__), '..'))

from aggregator import Aggregator
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

START = datetime(2026, 10, 19)
OUT = r"e:\MyProjects\maks_ahp\Week 48 Labour Turnover.docx"
DAYS = ["Mon", "Tue", "Wed", "Thur", "Fri", "Sat", "Sun"]


def pdate(s):
    c = re.sub(r"(\d+)(st|nd|rd|th)", r"\1", s or "", flags=re.I).replace(",", " ")
    for f in ("%d %B %Y", "%Y-%m-%d"):
        try:
            return datetime.strptime(re.sub(r"\s+", " ", c).strip(), f)
        except ValueError:
            pass


dailies = []
for f in glob.glob("cache/DAILY_*.json"):
    d = json.load(open(f, encoding="utf-8"))
    dt = pdate(d.get("date"))
    if dt and START <= dt <= START + timedelta(days=6):
        dailies.append(d)
assert len(dailies) == 7, f"expected 7 dailies, got {len(dailies)}"

data = asyncio.run(Aggregator(history_dir="cache/history_tmp").compile_weekly_data(dailies))
import shutil; shutil.rmtree("cache/history_tmp", ignore_errors=True)
labour = data["labour"]  # {category: [7 x (str | {"Day","Night"})]}


def dn(v):
    if isinstance(v, dict):
        return str(v.get("Day", "0")), str(v.get("Night", "0"))
    return str(v), None

# Expand a day into Day/Night columns only if any category has night data
has_night = [any(isinstance(vals[i], dict) for vals in labour.values()) for i in range(7)]
headers = ["CATEGORY"]
for i, d in enumerate(DAYS):
    if has_night[i]:
        headers += [f"{d} - Day", f"{d} - Night"]
    else:
        headers.append(d)

rows = []
for cat, vals in labour.items():
    row = [cat]
    for i in range(7):
        day, night = dn(vals[i])
        if has_night[i]:
            row += [day, night if night is not None else "0"]
        else:
            row.append(day)
    rows.append(row)

doc = Document()
sec = doc.sections[0]
sec.orientation = WD_ORIENT.LANDSCAPE
sec.page_width, sec.page_height = sec.page_height, sec.page_width
for m in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
    setattr(sec, m, Cm(1.2))

t = doc.add_paragraph()
t.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = t.add_run("MAKINDU AHP - LABOUR TURNOVER\nWeek 48: 19th - 25th October 2026")
r.bold = True; r.font.size = Pt(14)

table = doc.add_table(rows=1, cols=len(headers))
table.style = "Table Grid"


def shade(cell, color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), color)
    tcPr.append(shd)


def put(cell, text, bold=False, align=WD_ALIGN_PARAGRAPH.CENTER, color=None):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    run = p.add_run(text)
    run.font.size = Pt(7.5); run.bold = bold
    if color: run.font.color.rgb = RGBColor.from_string(color)

for i, h in enumerate(headers):
    put(table.rows[0].cells[i], h, True, color="FFFFFF")
    shade(table.rows[0].cells[i], "1F3864")

for row in rows:
    is_total = row[0].upper() == "TOTAL"
    cells = table.add_row().cells
    for i, v in enumerate(row):
        put(cells[i], v, bold=is_total or i == 0,
            align=WD_ALIGN_PARAGRAPH.LEFT if i == 0 else WD_ALIGN_PARAGRAPH.CENTER)
        if is_total:
            shade(cells[i], "D9E2F3")

doc.save(OUT)
print("saved", OUT, "| columns:", len(headers), "| rows:", len(rows))
for r_ in rows:
    if r_[0].upper() == "TOTAL": print(dict(zip(headers, r_)))
